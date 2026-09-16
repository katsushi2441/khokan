#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""地方厚生局「届出受理指定訪問看護事業所名簿」（Excel）を読み、厚労省CSVの事業所と突き合わせる。

  /usr/bin/python3 scripts/parse_kouseikyoku.py <名簿.xlsx …> [--db php/khokan_data/khokan.sqlite]

厚労省のCSV（介護保険の事業所番号）と厚生局の名簿（医療保険のステーションコード）には
共通のキーが無い。**電話番号**で突き合わせる（CSVは全軒に電話がある）。電話が合わない
ものは名前＋市区町村で補う。どちらでも決まらないものは結び付けない（推測しない）。

名簿の受理番号は、家族が知りたい「対応できるか」に直結する:
  訪看23/24 = 24時間対応体制加算（イ/ロ）  訪看10 = 精神科訪問看護基本療養費
  訪看25   = 特別管理加算                 訪看27/28 = 精神科の複数回訪問・重症患者支援
（略称表: 2606_houkan_ryakusyouhyou.pdf）
"""
import argparse
import re
import sqlite3
import sys
import unicodedata

import openpyxl

CODE = re.compile(r"\(\s*訪看(\d+)\s*\)")


def nfkc(s) -> str:
    return unicodedata.normalize("NFKC", str(s or "")).strip()


def tel_key(s) -> str:
    """先頭行の電話だけを数字にする（2行目はFAX）。"""
    first = nfkc(s).split("\n")[0]
    return re.sub(r"\D", "", first)


def name_key(s) -> str:
    s = nfkc(s)
    s = re.sub(r"(訪問看護リハビリステーション|訪問看護ステーション|訪問看護|ステーション|株式会社|合同会社|有限会社|医療法人社団|医療法人|社会福祉法人|一般社団法人|特定非営利活動法人)", "", s)
    return re.sub(r"[\s・()「」\-]", "", s)


WAREKI = {"令和": 2018, "平成": 1988, "昭和": 1925}


def asof_from(ws) -> str:
    """［令和 8年 8月 1日現在］を ISO 日付にする。見つからなければ空。"""
    for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        for c in r:
            m = re.search(r"(令和|平成|昭和)\s*(\d+)年\s*(\d+)月\s*(\d+)日現在", nfkc(c))
            if m:
                return f"{WAREKI[m.group(1)] + int(m.group(2)):04d}-{int(m.group(3)):02d}-{int(m.group(4)):02d}"
    return ""


def parse(path: str):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    parse.asof = asof_from(ws)
    out = []
    for r in ws.iter_rows(values_only=True):
        # 項番が数字の行だけが事業所（Excel 上は文字列のことがある）。見出し・頁ヘッダは飛ばす
        if len(r) < 19 or not str(r[2] if r[2] is not None else "").strip().isdigit():
            continue
        corp_name = nfkc(r[7]).split("\n")
        addr = nfkc(r[10]).split("\n")
        tel = nfkc(r[13]).split("\n")
        codes = sorted(set("訪看" + c for c in CODE.findall(nfkc(r[16]))))
        out.append({
            "seq": int(str(r[2]).strip()), "station_code": nfkc(r[4]),
            "corp": corp_name[0] if corp_name else "", "name": corp_name[1] if len(corp_name) > 1 else corp_name[0],
            "zip": addr[0] if addr else "", "address": "".join(addr[1:]) if len(addr) > 1 else "",
            "tel": tel[0] if tel else "", "tel_key": tel_key(r[13]),
            "codes": codes, "start": nfkc(r[18]).split("\n")[0],
        })
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--db", default="php/khokan_data/khokan.sqlite")
    ap.add_argument("--write", action="store_true", help="突き合わせ結果を DB の notices に書く（読むだけなら付けない）")
    ap.add_argument("--bureau", default="東海北陸厚生局", help="出典に書く厚生局名")
    ap.add_argument("--bureau-url", default="https://kouseikyoku.mhlw.go.jp/tokaihokuriku/newpage_00245.html")
    ap.add_argument("--asof", default="", help="名簿の時点（例 2026-08-01）。空なら Excel 内の［令和…現在］から読む")
    a = ap.parse_args()
    rows = []
    for f in a.files:
        rows.extend(parse(f))
    print(f"名簿: {len(rows)}軒（{len(a.files)}ファイル）")
    con = sqlite3.connect(f"file:{a.db}?mode=ro", uri=True)
    by_tel, by_name = {}, {}
    for no, name, city, tel in con.execute("SELECT no, name, city, tel FROM stations"):
        k = re.sub(r"\D", "", tel)
        if k:
            by_tel.setdefault(k, []).append(no)
        by_name.setdefault((name_key(name), city), []).append(no)
    hit_tel = hit_name = miss = 0
    samples = []
    for r in rows:
        nos = by_tel.get(r["tel_key"], [])
        if len(nos) == 1:
            hit_tel += 1; r["no"] = nos[0]; continue
        city = re.match(r"(.+?[市区町村])", r["address"])
        nos = by_name.get((name_key(r["name"]), city.group(1) if city else ""), [])
        if len(nos) == 1:
            hit_name += 1; r["no"] = nos[0]; continue
        miss += 1
        if len(samples) < 5:
            samples.append(r)
    tot = max(len(rows), 1)
    print(f"  電話で一致 {hit_tel}（{hit_tel/tot*100:.0f}%） 名前+市区町村で一致 {hit_name}（{hit_name/tot*100:.0f}%） 不一致 {miss}（{miss/tot*100:.0f}%）")
    c = {}
    for r in rows:
        for k in r["codes"]:
            c[k] = c.get(k, 0) + 1
    print("  受理番号の内訳:", {k: c[k] for k in sorted(c)})
    for s in samples:
        print("  不一致の例:", s["name"][:24], "|", s["address"][:22], "|", s["tel"])
    if not a.write:
        return 0
    asof = a.asof or getattr(parse, "asof", "")
    con.close()
    w = sqlite3.connect(a.db)
    w.executescript("""
    CREATE TABLE IF NOT EXISTS notices (
      no TEXT PRIMARY KEY,          -- 厚労省CSVの事業所番号（突き合わせ済みのものだけ）
      station_code TEXT, codes TEXT, start TEXT,
      bureau TEXT, bureau_url TEXT, asof TEXT
    );
    """)
    w.execute("DELETE FROM notices WHERE bureau=?", (a.bureau,))
    n = 0
    for r in rows:
        if "no" not in r:
            continue
        w.execute("INSERT OR REPLACE INTO notices VALUES (?,?,?,?,?,?,?)",
                  (r["no"], r["station_code"], ",".join(r["codes"]), r["start"], a.bureau, a.bureau_url, asof))
        n += 1
    w.execute("INSERT OR REPLACE INTO meta VALUES ('notices_" + re.sub(r"\W", "", a.bureau) + "', ?)", (f"{n}件 時点{asof}",))
    w.commit(); w.close()
    print(f"  notices に {n}件 書き込み（{a.bureau}・時点 {asof or '不明'}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
